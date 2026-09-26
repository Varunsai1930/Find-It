"""Disclosure intake: downloaded workbooks -> checked, parsed CSVs.

Synthetic workbooks in tmp folders only; no network, never ./tracker.db.
"""

from __future__ import annotations

import sqlite3
import zipfile
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd

import amfi_mf_parser
from findit.cli import intake as intake_cli
from findit.ingest import intake

HEADER = ["Name of the Instrument", "ISIN", "Industry", "Quantity",
          "Market Value (Rs. in Lakhs)", "% to NAV"]
HOLDINGS = [["Reliance Industries Ltd", "INE002A01018", "Oil", 100, 60.0, 60.0],
            ["Infosys Ltd", "INE009A01021", "IT", 50, 40.0, 40.0]]


def _book(path: Path, sheets: dict[str, str], as_on="August 31, 2026",
          header=HEADER, first_row=None) -> Path:
    """A workbook with one sheet per {sheet name: scheme title}."""
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, title in sheets.items():
        ws = wb.create_sheet(name)
        ws.append(first_row or [title])
        ws.append([f"Portfolio as on {as_on}"])
        ws.append([])
        ws.append(header)
        for row in HOLDINGS:
            ws.append(row)
    wb.save(path)
    return path


def _month(tmp_path: Path) -> Path:
    return tmp_path / "inbox" / "2026-08"


def test_accepted_folders_give_csvs_and_the_load_command(tmp_path):
    month = _month(tmp_path)
    _book(month / "Nippon India AMC" / "nippon.xlsx", {"GF": "Nippon India Growth Fund"})
    # A zip is unpacked and its workbooks parsed like loose ones.
    inner = _book(tmp_path / "dsp.xlsx", {"FLEXI": "DSP Flexi Cap Fund"})
    (month / "DSP AMC").mkdir(parents=True)
    with zipfile.ZipFile(month / "DSP AMC" / "dsp.zip", "w") as zf:
        zf.write(inner, "dsp.xlsx")

    report = intake.prepare(month, "2026-08")

    assert report.ok, intake.render(report)
    csvs = {r.amc: r.csv for r in report.results}
    for amc, sheet in (("Nippon India AMC", "GF"), ("DSP AMC", "FLEXI")):
        frame = pd.read_csv(csvs[amc])
        assert set(frame["amc_name"]) == {amc} and set(frame["report_month"]) == {"2026-08"}
        assert set(frame["scheme_name"]) == {sheet}
    command = intake.pipeline_command(report)
    assert "--prev 2026-07 --curr 2026-08" in command
    assert all(str(path) in command for path in csvs.values())
    missing = {a.amc: a.disclosure_page for a in report.missing}
    assert "Nippon India AMC" not in missing and len(missing) == 55
    assert missing["HDFC AMC"].startswith("https://www.hdfcfund.com/")


def test_misspelt_folder_is_not_parsed_and_suggests_the_registry_name(tmp_path):
    month = _month(tmp_path)
    _book(month / "Nipon India AMC" / "n.xlsx", {"GF": "Nippon India Growth Fund"})
    report = intake.prepare(month, "2026-08")
    assert not report.ok and report.results == []
    assert "Nippon India AMC" in report.unknown_folders["Nipon India AMC"]
    assert "did you mean 'Nippon India AMC'" in intake.render(report)


def test_the_same_sheet_in_two_files_refuses_the_amc_and_drops_its_old_csv(tmp_path):
    month = _month(tmp_path)
    folder = month / "SBI AMC"
    _book(folder / "all-schemes.xlsx", {"SBLUECHIP": "SBI Bluechip Fund"})
    first = intake.prepare(month, "2026-08")
    assert first.ok and first.results[0].csv.exists()
    # A per-scheme download on top of the consolidated one would load twice.
    _book(folder / "bluechip.xlsx", {"SBLUECHIP": "SBI Bluechip Fund"})
    second = intake.prepare(month, "2026-08")
    result = second.results[0]
    assert result.csv is None and not first.results[0].csv.exists()
    assert any("all-schemes.xlsx [SBLUECHIP] and bluechip.xlsx [SBLUECHIP]" in p
               for p in result.problems)
    assert intake.pipeline_command(second) is None


def test_sheets_the_loader_would_merge_count_as_one_scheme(tmp_path):
    month = _month(tmp_path)
    folder = month / "SBI AMC"
    _book(folder / "a.xlsx", {"SBI Bluechip Fund": "SBI Bluechip Fund"})
    _book(folder / "b.xlsx", {"SBI  BLUECHIP-Fund": "SBI Bluechip Fund"})
    result = intake.prepare(month, "2026-08").results[0]
    assert result.csv is None
    assert any("a.xlsx [SBI Bluechip Fund] and b.xlsx [SBI  BLUECHIP-Fund]" in p
               for p in result.problems)


def test_a_file_in_the_wrong_amc_folder_is_refused(tmp_path):
    month = _month(tmp_path)
    _book(month / "Nippon India AMC" / "n.xlsx",
          {"GF": "Nippon India Growth Fund", "SBLUECHIP": "SBI Bluechip Fund"})
    result = intake.prepare(month, "2026-08").results[0]
    assert result.csv is None
    assert any("a SBI AMC scheme: move its file to the SBI AMC folder" in p
               for p in result.problems)


def test_a_workbook_dated_another_month_is_refused(tmp_path):
    month = _month(tmp_path)
    _book(month / "SBI AMC" / "july.xlsx", {"SBLUECHIP": "SBI Bluechip Fund"},
          as_on="July 31, 2026")
    result = intake.prepare(month, "2026-08").results[0]
    assert result.problems == ["july.xlsx: dated 2026-07, not 2026-08"]
    # A header date written as a real Excel date counts too.
    path = _book(tmp_path / "dated.xlsx", {"S": "SBI Bluechip Fund"})
    wb = openpyxl.load_workbook(path)
    wb["S"]["B2"] = datetime(2026, 7, 31)
    wb["S"]["A2"] = "As on"
    wb.save(path)
    assert intake._months_named(path) == {"2026-07"}


def test_a_zip_escaping_its_folder_is_refused_and_nothing_escapes(tmp_path):
    month = _month(tmp_path)
    folder = month / "DSP AMC"
    folder.mkdir(parents=True)
    inner = _book(tmp_path / "x.xlsx", {"FLEXI": "DSP Flexi Cap Fund"})
    with zipfile.ZipFile(folder / "bad.zip", "w") as zf:
        zf.write(inner, "../../escaped.xlsx")
    result = intake.prepare(month, "2026-08").results[0]
    assert any("bad.zip: unsafe paths in zip" in p for p in result.problems)
    assert not (month / "escaped.xlsx").exists() and not (tmp_path / "escaped.xlsx").exists()


def test_an_unreadable_workbook_is_reported_while_other_amcs_still_parse(tmp_path):
    month = _month(tmp_path)
    no_value = [h for h in HEADER if not h.startswith("Market")]
    _book(month / "SBI AMC" / "broken.xlsx", {"S1": "SBI Bluechip Fund"}, header=no_value)
    _book(month / "DSP AMC" / "dsp.xlsx", {"FLEXI": "DSP Flexi Cap Fund"})
    report = intake.prepare(month, "2026-08")
    by_amc = {r.amc: r for r in report.results}
    assert by_amc["DSP AMC"].csv is not None
    assert by_amc["SBI AMC"].csv is None
    assert any(p.startswith("broken.xlsx: ValueError") and "market_value_lakhs" in p
               for p in by_amc["SBI AMC"].problems)
    assert intake_cli.main(["--month", "2026-08", "--inbox", str(month.parent)]) == 1


def test_excel_lock_files_are_ignored_and_loose_files_reported(tmp_path):
    month = _month(tmp_path)
    folder = month / "SBI AMC"
    _book(folder / "sbi.xlsx", {"SBLUECHIP": "SBI Bluechip Fund"})
    (folder / "~$sbi.xlsx").write_bytes(b"not a workbook")
    (month / "stray.xlsx").write_bytes(b"x")
    report = intake.prepare(month, "2026-08")
    assert report.results[0].files == ["sbi.xlsx"] and report.results[0].csv is not None
    assert report.loose_files == ["stray.xlsx"] and not report.ok


def test_cli_reads_the_database_read_only_and_never_creates_it(tmp_path):
    inbox = tmp_path / "inbox"
    missing_db = tmp_path / "absent.db"
    assert intake_cli.main(["--month", "2026-08", "--inbox", str(inbox),
                            "--db", str(missing_db), "--init"]) == 0
    assert not missing_db.exists() and (inbox / "2026-08").is_dir()

    db_path = tmp_path / "t.db"
    conn = sqlite3.connect(db_path)
    conn.execute("CREATE TABLE schemes (scheme_id INTEGER, amc_name TEXT)")
    conn.executemany("INSERT INTO schemes VALUES (?, ?)", [(1, "SBI AMC"), (2, "Made Up AMC")])
    conn.commit()
    conn.close()
    before = db_path.read_bytes()
    assert intake_cli.main(["--month", "2026-09", "--inbox", str(inbox),
                            "--db", str(db_path), "--init"]) == 0
    # Only AMCs AMFI lists get a folder; the database is unchanged.
    assert sorted(p.name for p in (inbox / "2026-09").iterdir()) == ["SBI AMC"]
    assert db_path.read_bytes() == before
    assert intake_cli.main(["--month", "2026-13", "--inbox", str(inbox)]) == 2


def test_registry_names_match_the_database_convention():
    registry = intake.load_registry()
    names = [a.amc for a in registry]
    brands = [b for a in registry for b in a.brands]
    assert len(registry) == 57 and len(set(names)) == 57 and len(set(brands)) == len(brands)
    # The names the database already uses for the AMCs loaded by hand.
    assert {"SBI AMC", "HDFC AMC", "ICICI Prudential AMC"} <= set(names)
    assert intake._brand_owner("Parag Parikh Flexi Cap Fund", registry) == "PPFAS AMC"
    # Whole words only: a "quant" brand must not claim Quantum's schemes.
    assert intake._brand_owner("Quantum Value Fund", registry) == "Quantum AMC"
    assert intake._brand_owner("Bharat 22 ETF", registry) is None


def test_previous_month_wraps_the_year():
    assert intake.previous_month("2026-01") == "2025-12"
    assert intake.previous_month("2026-08") == "2026-07"


# -- parser: layouts from AMCs loaded through the intake -------------------------


def test_parser_reads_nippon_style_title_rows_and_headers(tmp_path):
    # Nippon: scheme code, then the title, then a back-link cell reading "Index".
    header = ["ISIN", "Name of the Instrument", "Industry / Rating", "Quantity",
              "Market/Fair Value\n( Rs. in Lacs)", "% to NAV"]
    path = _book(tmp_path / "nippon.xlsx", {"GF": ""}, header=header,
                 first_row=["RLMF001", "Nippon India Growth Mid Cap Fund (An open-ended "
                            "equity scheme)", "Index"])
    wb = openpyxl.load_workbook(path)
    for row in wb["GF"].iter_rows(min_row=5):
        name, isin = row[0].value, row[1].value
        row[0].value, row[1].value = isin, name  # ISIN first, as Nippon lays it out
    wb.save(path)
    frame = amfi_mf_parser.parse_workbook(path, "Nippon India AMC", "2026-08")
    assert set(frame["scheme_title"]) == {
        "Nippon India Growth Mid Cap Fund (An open-ended equity scheme)"}
    assert frame["market_value_lakhs"].tolist() == [60.0, 40.0]
